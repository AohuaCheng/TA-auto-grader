#!/usr/bin/env python3
"""Export Essay 1 first-round scores to an xlsx summary (template-based)."""

from __future__ import annotations

import argparse
import ast
import re
import shutil
from pathlib import Path

import openpyxl
from docx import Document

from paths import download_root, templates_dir
from topic_detector import build_topic_field_from_docx
from yuketang_grader import collect_docx_jobs, feedback_path_for

TEMPLATE_NAME = "MINC 2025 Essay 分数汇总 样例.xlsx"
DEFAULT_OUTPUT_NAME = "Essay1 分数汇总.xlsx"
SHEET_NAME = "Essay 1"
FIRST_ROUND_END_COL = 15  # A-O in template (before second submission block)

SCORE_FIELDS = (
    ("Ideational Flexibility", "diversity"),
    ("Originality", "originality"),
    ("Relevance", "relevance"),
    ("Evaluation", "evaluation"),
    ("Analysis", "analysis"),
    ("Inference", "inference"),
    ("Explanation", "explanation"),
    ("Interpretation", "interpretation"),
    ("Factuality", "factuality"),
    ("Clarity", "clarity"),
)


def template_path() -> Path:
    path = templates_dir() / TEMPLATE_NAME
    if not path.exists():
        raise FileNotFoundError(f"找不到模板: {path}")
    return path


def topic_label(docx_path: Path) -> str:
    topic_id, _ = build_topic_field_from_docx(docx_path)
    return "Plants" if topic_id == 1 else "Psyche"


def extract_scores(feedback_path: Path) -> dict[str, float]:
    doc = Document(str(feedback_path))
    for paragraph in reversed(doc.paragraphs):
        text = paragraph.text
        if "Overall" in text and "Ideational Flexibility" in text:
            match = re.search(r"\{.*\}", text, re.DOTALL)
            if match:
                data = ast.literal_eval(match.group())
                return {str(k): float(v) for k, v in data.items()}
    raise ValueError(f"未找到分数: {feedback_path}")


def prepare_workbook(template: Path, output: Path) -> openpyxl.Workbook:
    shutil.copy2(template, output)
    wb = openpyxl.load_workbook(output)

    for name in list(wb.sheetnames):
        if name != SHEET_NAME:
            del wb[name]

    ws = wb[SHEET_NAME]

    if ws.max_column > FIRST_ROUND_END_COL:
        ws.delete_cols(FIRST_ROUND_END_COL + 1, ws.max_column - FIRST_ROUND_END_COL)

    if ws.max_row > 2:
        ws.delete_rows(3, ws.max_row - 2)

    for merged in list(ws.merged_cells.ranges):
        try:
            ws.unmerge_cells(str(merged))
        except KeyError:
            pass

    ws.insert_cols(3)

    ws.cell(row=1, column=1, value="学号")
    ws.cell(row=1, column=2, value="姓名")
    ws.cell(row=1, column=3, value="Topic")
    ws.cell(row=1, column=7, value="Sum")
    ws.cell(row=1, column=13, value="Sum")
    ws.merge_cells("D1:F1")
    ws["D1"] = "Creativity"
    ws.merge_cells("H1:L1")
    ws["H1"] = "Critical Thinking"

    ws.cell(row=2, column=4, value="Diversity")
    ws.cell(row=2, column=5, value="Originality")
    ws.cell(row=2, column=6, value="Relevance")
    ws.cell(row=2, column=7, value="Out of 45")
    ws.cell(row=2, column=8, value="Evaluation")
    ws.cell(row=2, column=9, value="Analysis ")
    ws.cell(row=2, column=10, value="Inference")
    ws.cell(row=2, column=11, value="Explanation")
    ws.cell(row=2, column=12, value="Interpretation")
    ws.cell(row=2, column=13, value="Out of 45")
    ws.cell(row=2, column=14, value="Factuality")
    ws.cell(row=2, column=15, value="Clarity")
    ws.cell(row=2, column=16, value="Overall")

    return wb


def write_row(ws, row_idx: int, student_id: str, name: str, topic: str, scores: dict[str, float]) -> None:
    ws.cell(row=row_idx, column=1, value=student_id)
    ws.cell(row=row_idx, column=2, value=name)
    ws.cell(row=row_idx, column=3, value=topic)

    # D-F: creativity block (Diversity = Ideational Flexibility)
    ws.cell(row=row_idx, column=4, value=scores.get("Ideational Flexibility"))
    ws.cell(row=row_idx, column=5, value=scores.get("Originality"))
    ws.cell(row=row_idx, column=6, value=scores.get("Relevance"))
    ws.cell(row=row_idx, column=7, value=f"=SUM(D{row_idx}:F{row_idx})")

    # H-L: critical thinking block
    ws.cell(row=row_idx, column=8, value=scores.get("Evaluation"))
    ws.cell(row=row_idx, column=9, value=scores.get("Analysis"))
    ws.cell(row=row_idx, column=10, value=scores.get("Inference"))
    ws.cell(row=row_idx, column=11, value=scores.get("Explanation"))
    ws.cell(row=row_idx, column=12, value=scores.get("Interpretation"))
    ws.cell(row=row_idx, column=13, value=f"=SUM(H{row_idx}:L{row_idx})")

    ws.cell(row=row_idx, column=14, value=scores.get("Factuality"))
    ws.cell(row=row_idx, column=15, value=scores.get("Clarity"))
    ws.cell(row=row_idx, column=16, value=scores.get("Overall"))


def export_score_summary(
    download_dir: Path,
    *,
    template: Path | None = None,
    output_path: Path | None = None,
) -> Path:
    src_template = template or template_path()
    target = output_path or (download_dir / DEFAULT_OUTPUT_NAME)

    jobs = sorted(collect_docx_jobs(download_dir), key=lambda j: j.student_id)
    if not jobs:
        raise ValueError(f"未找到学生作业: {download_dir}")

    wb = prepare_workbook(src_template, target)
    ws = wb[SHEET_NAME]

    for idx, job in enumerate(jobs, start=3):
        feedback = feedback_path_for(job)
        scores = extract_scores(feedback)
        label = topic_label(job.docx_path)
        write_row(ws, idx, job.student_id, job.name, label, scores)

    wb.save(target)
    return target


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="导出 Essay 1 第一次批改分数汇总 xlsx")
    parser.add_argument("--dir", help="作业目录（默认 downloads 下唯一子目录）")
    parser.add_argument("--template", help="分数汇总模板 xlsx")
    parser.add_argument("--output", help="输出文件路径")
    args = parser.parse_args(argv)

    if args.dir:
        download_dir = Path(args.dir)
    else:
        root = download_root()
        subdirs = sorted(p for p in root.iterdir() if p.is_dir())
        download_dir = subdirs[0] if len(subdirs) == 1 else root

    output = Path(args.output) if args.output else None
    template = Path(args.template) if args.template else None
    path = export_score_summary(download_dir, template=template, output_path=output)
    print(f"已生成: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
