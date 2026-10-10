#!/usr/bin/env python3
"""Extract Review Comments from Yuketang feedback and save upload-ready docx."""

from __future__ import annotations

import argparse
import re
import shutil
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt

from paths import download_root, templates_dir
from yuketang_grader import collect_docx_jobs, feedback_path_for

UPLOAD_TEMPLATE_NAME = "仅反馈样例 这样上传至网络学堂.docx"
OUTPUT_NAME_TEMPLATE = "Firstround_feedback_{student_id}.docx"

REVIEW_MARKER = re.compile(
    r"Review Comments(?:\s+for\s+Assignment)?",
    re.IGNORECASE,
)
HEADER_LINE_RE = re.compile(r"^#{1,6}\s+")
RULE_LINE_RE = re.compile(r"^-{3,}\s*$")
SECTION1_RE = re.compile(r"^1\.\s*Main\s+content\s+and\s+advantages", re.I)
SECTION2_RE = re.compile(r"^2\.\s*Areas?\s+of\s+deficiency\s+and\s+suggestions?", re.I)
CATEGORY_NAMES = ("Critical Thinking", "Creativity", "Clarity and Factuality")
CATEGORY_LINE_RE = re.compile(
    r"^-\s*(Critical Thinking|Creativity|Clarity and Factuality)\s*:\s*$",
    re.I,
)
CATEGORY_HEADER_RE = re.compile(
    r"^(Critical Thinking|Creativity|Clarity and Factuality)\s*:?\s*$",
    re.I,
)
KEY_DEFICIENCY_RE = re.compile(r"^Key Deficiency\s+(\d+)\s*:\s*(.*)$", re.I)
BULLET_LINE_RE = re.compile(r"^-\s*(?:Sub-bullet:\s*)?(.*)$", re.I)
INDENT_BULLET_RE = re.compile(r"^(\s+)-\s*(?:Sub-bullet:\s*)?(.*)$", re.I)


@dataclass
class ContentBlock:
    kind: str
    text: str
    level: int = 0
    number: int | None = None


def upload_template_path() -> Path:
    path = templates_dir() / UPLOAD_TEMPLATE_NAME
    if not path.exists():
        raise FileNotFoundError(f"找不到上传模板: {path}")
    return path


def output_path_for(student_id: str, student_dir: Path) -> Path:
    return student_dir / OUTPUT_NAME_TEMPLATE.format(student_id=student_id)


def _clear_document_body(doc: Document) -> None:
    body = doc.element.body
    for child in list(body):
        if child.tag.split("}")[-1] in ("p", "tbl"):
            body.remove(child)


def extract_review_comments_text(feedback_path: Path) -> str:
    doc = Document(str(feedback_path))
    for paragraph in reversed(doc.paragraphs):
        if REVIEW_MARKER.search(paragraph.text):
            return paragraph.text
    raise ValueError(f"未找到 Review Comments 段落: {feedback_path}")


def slice_review_body(full_text: str) -> str:
    match = REVIEW_MARKER.search(full_text)
    if not match:
        raise ValueError("Review Comments 标记未找到")

    after = full_text[match.end() :]
    after = re.sub(r'^\s*:\s*"[^"]*"\s*', "", after, count=1)
    after = after.lstrip("\n")

    lines = after.splitlines()
    cleaned: list[str] = []
    started = False
    for line in lines:
        stripped = line.strip()
        if not started:
            if RULE_LINE_RE.match(stripped) or not stripped:
                continue
            if re.search(r"main content and advantages", stripped, re.I):
                started = True
                cleaned.append(line)
                continue
            if HEADER_LINE_RE.match(stripped):
                continue
            continue
        cleaned.append(line)

    if not cleaned:
        raise ValueError("Review Comments 正文为空")
    return "\n".join(cleaned)


def normalize_category(name: str) -> str:
    key = name.strip().rstrip(":").lower()
    mapping = {
        "critical thinking": "Critical Thinking:",
        "creativity": "Creativity:",
        "clarity and factuality": "Clarity and Factuality:",
    }
    return mapping.get(key, name.strip() + ("" if name.strip().endswith(":") else ":"))


def parse_inline_segments(text: str) -> list[tuple[str, bool]]:
    """Split text into (segment, is_italic) preserving *italic* markers."""
    segments: list[tuple[str, bool]] = []
    pos = 0
    for match in re.finditer(r"\*([^*]+)\*", text):
        if match.start() > pos:
            segments.append((text[pos : match.start()], False))
        segments.append((match.group(1), True))
        pos = match.end()
    if pos < len(text):
        segments.append((text[pos:], False))
    if not segments:
        segments.append((text, False))
    return [(re.sub(r"\s+", " ", s).strip(), italic) for s, italic in segments if s.strip()]


def parse_review_comments(markdown_text: str) -> list[ContentBlock]:
    body = slice_review_body(markdown_text)
    blocks: list[ContentBlock] = []
    in_section2 = False

    for raw_line in body.splitlines():
        line = raw_line.rstrip("\n")
        stripped = line.strip()
        if not stripped or RULE_LINE_RE.match(stripped):
            continue

        if HEADER_LINE_RE.match(stripped):
            title = HEADER_LINE_RE.sub("", stripped).strip()
            title = re.sub(r"\s+", " ", title)
            if SECTION1_RE.match(title):
                in_section2 = False
                blocks.append(ContentBlock("section", "1. Main content and advantages."))
                continue
            if SECTION2_RE.match(title):
                in_section2 = True
                blocks.append(
                    ContentBlock("section", "2. Areas of deficiency and suggestions.")
                )
                continue
            cat = CATEGORY_HEADER_RE.match(title)
            if cat:
                blocks.append(ContentBlock("category", normalize_category(cat.group(1))))
                continue
            continue

        cat = CATEGORY_LINE_RE.match(stripped)
        if cat:
            blocks.append(ContentBlock("category", normalize_category(cat.group(1))))
            continue

        key_def = KEY_DEFICIENCY_RE.match(stripped)
        if key_def and in_section2:
            blocks.append(
                ContentBlock(
                    "deficiency",
                    key_def.group(2).strip(),
                    number=int(key_def.group(1)),
                )
            )
            continue

        indent_match = INDENT_BULLET_RE.match(line)
        if indent_match:
            content = indent_match.group(2).strip()
            if in_section2:
                blocks.append(ContentBlock("suggestion", content, level=1))
            else:
                blocks.append(ContentBlock("bullet", content, level=1))
            continue

        bullet = BULLET_LINE_RE.match(stripped)
        if bullet:
            content = bullet.group(1).strip()
            if in_section2:
                blocks.append(ContentBlock("suggestion", content, level=0))
            else:
                blocks.append(ContentBlock("bullet", content, level=0))
            continue

        plain = stripped.lstrip("-").strip()
        if plain:
            blocks.append(ContentBlock("body", plain))

    return blocks


def _indent_for_level(level: int) -> Pt:
    return Pt(18 + level * 18)


def _apply_readable_style(paragraph, *, bold: bool = False, level: int = 0) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = 1.35
    fmt.space_after = Pt(6)
    if level:
        fmt.left_indent = _indent_for_level(level)
        fmt.first_line_indent = Pt(-12)


def _add_rich_text(paragraph, text: str, *, prefix: str = "", bold_prefix: bool = False) -> None:
    if prefix:
        run = paragraph.add_run(prefix)
        run.bold = bold_prefix
    segments = parse_inline_segments(text)
    if not segments:
        paragraph.add_run(text)
        return
    for segment, italic in segments:
        run = paragraph.add_run(segment)
        run.italic = italic


def build_upload_document(blocks: list[ContentBlock], template_path: Path) -> Document:
    doc = Document(str(template_path))
    _clear_document_body(doc)

    for block in blocks:
        if not block.text.strip() and block.kind != "section":
            continue

        if block.kind == "section":
            para = doc.add_paragraph(style="Normal")
            _apply_readable_style(para, bold=True)
            run = para.add_run(block.text)
            run.bold = True
            continue

        if block.kind == "category":
            para = doc.add_paragraph(style="Normal")
            _apply_readable_style(para)
            run = para.add_run(block.text)
            run.bold = True
            continue

        if block.kind == "bullet":
            para = doc.add_paragraph(style="Normal")
            _apply_readable_style(para, level=block.level)
            _add_rich_text(para, block.text, prefix="• ")
            continue

        if block.kind == "deficiency":
            para = doc.add_paragraph(style="Normal")
            _apply_readable_style(para, level=0)
            label = f"Deficiency {block.number}: " if block.number else "Deficiency: "
            _add_rich_text(para, block.text, prefix=label, bold_prefix=True)
            continue

        if block.kind == "suggestion":
            para = doc.add_paragraph(style="Normal")
            _apply_readable_style(para, level=block.level + 1)
            _add_rich_text(para, block.text, prefix="• Suggestion: ")
            continue

        para = doc.add_paragraph(style="Normal")
        _apply_readable_style(para, level=block.level)
        _add_rich_text(para, block.text)

    return doc


def convert_feedback_to_firstround(
    feedback_path: Path,
    student_id: str,
    *,
    template_path: Path | None = None,
    output_path: Path | None = None,
) -> Path:
    template = template_path or upload_template_path()
    target = output_path or output_path_for(student_id, feedback_path.parent)

    markdown = extract_review_comments_text(feedback_path)
    blocks = parse_review_comments(markdown)
    if not blocks:
        raise ValueError(f"Review Comments 解析结果为空: {feedback_path}")

    shutil.copy2(template, target)
    doc = build_upload_document(blocks, template)
    doc.save(str(target))
    return target


def run_format_feedback(
    download_dir: Path,
    *,
    template_path: Path | None = None,
    student: str = "",
    dry_run: bool = False,
) -> list[tuple[Path, str]]:
    template = template_path or upload_template_path()
    results: list[tuple[Path, str]] = []

    jobs = collect_docx_jobs(download_dir)
    if student:
        needle = student.strip().lower()
        jobs = [
            j
            for j in jobs
            if needle in j.student_id.lower() or needle in j.name.lower()
        ]

    for job in jobs:
        fb_path = feedback_path_for(job)
        out_path = output_path_for(job.student_id, fb_path.parent)
        if not fb_path.exists():
            results.append((out_path, "跳过（无 feedback）"))
            continue

        try:
            markdown = extract_review_comments_text(fb_path)
            blocks = parse_review_comments(markdown)
            kinds = {}
            for b in blocks:
                kinds[b.kind] = kinds.get(b.kind, 0) + 1
            summary = (
                f"预览: {out_path.name} "
                f"({', '.join(f'{k}={v}' for k, v in sorted(kinds.items()))})"
            )
            if dry_run:
                results.append((out_path, summary))
                continue

            convert_feedback_to_firstround(
                fb_path,
                job.student_id,
                template_path=template,
                output_path=out_path,
            )
            results.append((out_path, "已生成"))
        except (ValueError, OSError) as exc:
            results.append((out_path, f"失败: {exc}"))

    return results


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="从 feedback 提取 Review Comments，生成 Firstround_feedback_学号.docx"
    )
    parser.add_argument("--dir", help="作业目录（默认 downloads 下唯一子目录）")
    parser.add_argument(
        "--template",
        help=f"上传模板（默认 templates/{UPLOAD_TEMPLATE_NAME}）",
    )
    parser.add_argument("--student", help="仅处理指定学号/姓名")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args(argv)

    if args.dir:
        download_dir = Path(args.dir)
    else:
        root = download_root()
        subdirs = sorted(p for p in root.iterdir() if p.is_dir())
        download_dir = subdirs[0] if len(subdirs) == 1 else root

    template = Path(args.template) if args.template else None
    results = run_format_feedback(
        download_dir,
        template_path=template,
        student=args.student or "",
        dry_run=args.dry_run,
    )

    ok = sum(1 for _, status in results if status.startswith(("已生成", "预览")))
    print(f"目录: {download_dir}")
    print(f"处理: {ok}/{len(results)}")
    for path, status in results:
        print(f"  {path.parent.name}/{path.name}: {status}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
