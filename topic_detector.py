"""Detect which Essay 1 topic a student chose from their docx."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document

TOPIC_1_TITLE = "Do plants have consciousness?"
TOPIC_2_TITLE = (
    'Find out more about "psyche", and share your understanding to '
    "its symbolic meaning."
)
REQUIREMENT_TEXT = (
    "Show critical and creative thinking by engaging with multiple perspectives, "
    "including counterarguments. Balance academic evidence with personal insights, "
    "using critical thinking to question assumptions and creativity to develop "
    "fresh interpretations."
)


def extract_docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def detect_topic(text: str) -> int:
    lowered = text.lower()

    topic1_signals = [
        "do plants have consciousness" in lowered,
        bool(re.search(r"\bplant(s)?\b", lowered) and "conscious" in lowered),
    ]
    topic2_signals = [
        "psyche" in lowered,
        "symbolic meaning" in lowered,
        bool(re.search(r"\bpsyche\b", lowered) and "symbol" in lowered),
    ]

    score1 = sum(topic1_signals) + len(re.findall(r"conscious|plant", lowered)) * 0.1
    score2 = sum(topic2_signals) + len(re.findall(r"psyche|symbolic", lowered)) * 0.1

    if score1 == score2 == 0:
        return 1
    return 1 if score1 >= score2 else 2


def build_topic_field(text: str) -> tuple[int, str]:
    topic_id = detect_topic(text)
    title = TOPIC_1_TITLE if topic_id == 1 else TOPIC_2_TITLE
    return topic_id, f"{title} {REQUIREMENT_TEXT}"


def build_topic_field_from_docx(path: Path) -> tuple[int, str]:
    return build_topic_field(extract_docx_text(path))
