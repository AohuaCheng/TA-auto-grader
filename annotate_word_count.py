#!/usr/bin/env python3
"""Annotate score summary xlsx with essay word counts and out-of-range remarks."""

from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl

from paths import download_root
from word_count import (
    MAX_WORDS,
    MIN_WORDS,
    count_essay_words,
    merge_remarks,
    resolve_submission_docx,
    strip_under_word_remarks,
    word_count_remark,
)

DEFAULT_SCORE_XLSX = "Minc2026Essays分数汇总.xlsx"
DEFAULT_ASSIGNMENT_DIR = "Essay1"
SHEET_NAME = "Essay 1"
COL_WORD_COUNT = 18
COL_REMARK = 19


def annotate_word_count(
    *,
    score_xlsx: Path,
    assignment_dir: Path,
    min_words: int = MIN_WORDS,
    max_words: int = MAX_WORDS,
) -> Path:
    wb = openpyxl.load_workbook(score_xlsx)
    ws = wb[SHEET_NAME]

    if ws.cell(row=1, column=COL_WORD_COUNT).value == "备注":
        ws.insert_cols(COL_WORD_COUNT)

    ws.cell(row=1, column=COL_WORD_COUNT, value="字数")
    ws.cell(row=1, column=COL_REMARK, value="备注")

    below: list[str] = []
    above: list[str] = []

    for row_idx in range(3, ws.max_row + 1):
        student_id = ws.cell(row_idx, 1).value
        name = ws.cell(row_idx, 2).value
        if not student_id:
            continue

        student_dir = assignment_dir / f"{student_id}_{name}"
        docx_path = resolve_submission_docx(student_dir)
        if not docx_path:
            ws.cell(row_idx, COL_WORD_COUNT, value=None)
            continue

        word_count = count_essay_words(docx_path)
        ws.cell(row_idx, COL_WORD_COUNT, value=word_count)

        existing_remark = strip_under_word_remarks(ws.cell(row_idx, COL_REMARK).value)
        extra = word_count_remark(
            word_count,
            min_words=min_words,
            max_words=max_words,
            flag_under=False,
        )
        if existing_remark and extra and "字数超出" in extra and "字数超出" in existing_remark:
            ws.cell(row_idx, COL_REMARK, value=existing_remark)
        else:
            ws.cell(row_idx, COL_REMARK, value=merge_remarks(existing_remark, extra))

        label = f"{student_id} {name} ({word_count})"
        if extra and "不足" in extra:
            below.append(label)
        elif extra and "超出" in extra:
            above.append(label)

    wb.save(score_xlsx)
    print(f"已更新: {score_xlsx}")
    print(f"字数范围 {min_words}-{max_words}（不含题目与参考文献）")
    print(f"字数不足: {len(below)} 人")
    for item in below:
        print(f"  - {item}")
    print(f"字数超出: {len(above)} 人")
    for item in above:
        print(f"  - {item}")
    return score_xlsx


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="在分数汇总 xlsx 中标注论文字数")
    parser.add_argument("--scores", type=Path, help=f"分数汇总 xlsx（默认 downloads/{DEFAULT_SCORE_XLSX}）")
    parser.add_argument("--dir", type=Path, help=f"作业目录（默认 downloads/{DEFAULT_ASSIGNMENT_DIR}）")
    parser.add_argument("--min-words", type=int, default=MIN_WORDS)
    parser.add_argument("--max-words", type=int, default=MAX_WORDS)
    args = parser.parse_args(argv)

    root = download_root()
    score_xlsx = args.scores or (root / DEFAULT_SCORE_XLSX)
    assignment_dir = args.dir or (root / DEFAULT_ASSIGNMENT_DIR)

    annotate_word_count(
        score_xlsx=score_xlsx,
        assignment_dir=assignment_dir,
        min_words=args.min_words,
        max_words=args.max_words,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
