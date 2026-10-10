#!/usr/bin/env python3
"""Build learn.tsinghua.edu.cn grade-upload CSV from score summary + student folders."""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

import openpyxl

from paths import download_root

DEFAULT_SCORE_XLSX = "Minc2026Essays分数汇总.xlsx"
DEFAULT_ASSIGNMENT_DIR = "Essay1"
FIRST_ROUND_TEMPLATE = "Firstround_feedback_{student_id}.docx"

DEFAULT_COMMENT = (
    "Excellent work! Your essay demonstrates strong critical and creative thinking. "
    "Keep up the great effort. Please see the attached file for detailed feedback."
)

OFF_TOPIC_COMMENT = (
    "Your submission did not address the assigned topic, so it could not be graded "
    "on the essay requirements. Please read the prompt carefully and respond to the "
    "specific question next time. Please see the attached file for detailed feedback."
)

REMARK_TRANSLATIONS = (
    (
        re.compile(r"提交格式非 Word（原文件：(.+?)），请下次提交 \.doc 或 \.docx"),
        r"Note: Your submission was not in Word format (original file: \1). "
        r"Please submit a .doc or .docx file next time.",
    ),
    (
        re.compile(r"作业内容与所选题目无关，请下次根据题目要求作答"),
        "Note: Your submission did not address the assigned topic. "
        "Please read the prompt carefully and respond to the specific question next time.",
    ),
)


def translate_remark(remark: str) -> str:
    text = remark.strip()
    for pattern, replacement in REMARK_TRANSLATIONS:
        text = pattern.sub(replacement, text)
    return text


def build_comment(remark: str | None, *, score: float | int | None = None) -> str:
    remark_text = str(remark).strip() if remark else ""
    if "题目无关" in remark_text or score == 0:
        return OFF_TOPIC_COMMENT
    if remark_text:
        return f"{DEFAULT_COMMENT}\n\n{translate_remark(remark_text)}"
    return DEFAULT_COMMENT


def load_submissions(submissions_csv: Path) -> dict[str, dict[str, str]]:
    mapping: dict[str, dict[str, str]] = {}
    with submissions_csv.open(encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            sid = str(row.get("学号", "")).strip()
            if sid:
                mapping[sid] = row
    return mapping


def find_firstround_attachment(student_dir: Path, student_id: str) -> Path | None:
    expected = student_dir / FIRST_ROUND_TEMPLATE.format(student_id=student_id)
    if expected.exists():
        return expected
    matches = sorted(student_dir.glob("Firstround_feedback_*.docx"))
    return matches[0] if matches else None


def prepare_upload_csv(
    *,
    score_xlsx: Path,
    assignment_dir: Path,
    output_csv: Path,
    default_wlkcid: str | None = None,
) -> Path:
    submissions_csv = assignment_dir / "submissions.csv"
    if not submissions_csv.exists():
        raise FileNotFoundError(f"找不到 submissions.csv: {submissions_csv}")

    submissions = load_submissions(submissions_csv)
    wb = openpyxl.load_workbook(score_xlsx, data_only=True)
    ws = wb["Essay 1"]

    rows: list[dict[str, str]] = []
    missing_attachment: list[str] = []
    missing_xszyid: list[str] = []

    for r in range(3, ws.max_row + 1):
        student_id = ws.cell(r, 1).value
        if not student_id:
            continue
        student_id = str(student_id).strip()
        name = str(ws.cell(r, 2).value or "").strip()
        overall = ws.cell(r, 16).value
        excellent = ws.cell(r, 17).value
        remark = ws.cell(r, 18).value

        student_dir = assignment_dir / f"{student_id}_{name}"
        attachment = find_firstround_attachment(student_dir, student_id)
        if not attachment:
            missing_attachment.append(f"{student_id}_{name}")
            continue

        sub = submissions.get(student_id, {})
        xszyid = str(sub.get("xszyid", "")).strip()
        if not xszyid:
            missing_xszyid.append(f"{student_id}_{name}")
            continue

        meta_wlkcid = default_wlkcid
        meta_path = student_dir / "meta.json"
        if meta_path.exists() and not meta_wlkcid:
            import json

            meta = json.loads(meta_path.read_text(encoding="utf-8"))
            meta_wlkcid = meta.get("wlkcid")

        rows.append(
            {
                "学号": student_id,
                "姓名": name,
                "xszyid": xszyid,
                "wlkcid": meta_wlkcid or "",
                "成绩": str(int(overall)) if isinstance(overall, (int, float)) else str(overall or ""),
                "评语": build_comment(
                    str(remark) if remark else None,
                    score=overall if isinstance(overall, (int, float)) else None,
                ),
                "附件": str(attachment.resolve()),
                "优秀作业": "1" if excellent == 1 else "",
            }
        )

    if missing_attachment:
        raise FileNotFoundError(
            "以下学生缺少 Firstround_feedback 附件:\n" + "\n".join(missing_attachment)
        )
    if missing_xszyid:
        raise ValueError(
            "以下学生在 submissions.csv 中缺少 xszyid:\n" + "\n".join(missing_xszyid)
        )

    output_csv.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = ["学号", "姓名", "xszyid", "wlkcid", "成绩", "评语", "附件", "优秀作业"]
    with output_csv.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    excellent_count = sum(1 for row in rows if row["优秀作业"] == "1")
    print(f"已生成上传 CSV: {output_csv} ({len(rows)} 人，优秀作业 {excellent_count} 人)")
    return output_csv


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="生成绩上传 CSV（网络学堂批阅）")
    parser.add_argument(
        "--scores",
        type=Path,
        help=f"分数汇总 xlsx（默认 downloads/{DEFAULT_SCORE_XLSX}）",
    )
    parser.add_argument(
        "--dir",
        type=Path,
        help=f"作业目录（默认 downloads/{DEFAULT_ASSIGNMENT_DIR}）",
    )
    parser.add_argument(
        "--output",
        type=Path,
        help="输出 CSV 路径（默认作业目录/learn_upload.csv）",
    )
    parser.add_argument("--wlkcid", help="默认 wlkcid")
    args = parser.parse_args(argv)

    root = download_root()
    score_xlsx = args.scores or (root / DEFAULT_SCORE_XLSX)
    assignment_dir = args.dir or (root / DEFAULT_ASSIGNMENT_DIR)
    output_csv = args.output or (assignment_dir / "learn_upload.csv")

    prepare_upload_csv(
        score_xlsx=score_xlsx,
        assignment_dir=assignment_dir,
        output_csv=output_csv,
        default_wlkcid=args.wlkcid,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
